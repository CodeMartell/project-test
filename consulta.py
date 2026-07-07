import openpyxl


def consultar_usuario():
    print("=" * 45)
    print("🤖 SISTEMA DE CONSULTA DE USUÁRIOS — RPA")
    print("=" * 45)

    nome_arquivo = "Dados do Roteiro 06.xlsx"

    try:
        # Abre o arquivo Excel de forma nativa
        wb = openpyxl.load_workbook(nome_arquivo, data_only=True)
        # Seleciona a primeira planilha ativa
        aba = wb.active
    except FileNotFoundError:
        print(f"\n❌ Erro de Validação: O arquivo '{nome_arquivo}' não foi encontrado.")
        return

    # Entrada para pesquisa de usuários
    pesquisa = input("Digite o Nome ou CPF para pesquisar: ").strip().lower()
    encontrado = False

    # Percorre as linhas do Excel (começando da linha 2 para pular o cabeçalho)
    for linha in aba.iter_rows(min_row=2, max_row=aba.max_row, values_only=True):
        # Se a linha estiver vazia, pula
        if not linha[0]:
            continue

        # Mapeamento das colunas com base na foto do roteiro
        nome = str(linha[0]).strip()
        sobrenome = str(linha[1]).strip()
        cpf = str(linha[2]).strip()
        email = str(linha[3]).strip()
        telefone = str(linha[4]).strip()
        nascimento = str(linha[5]).strip()
        status = str(linha[6]).strip()
        observacao = str(linha[7]).strip()
        endereco = str(linha[8]).strip()

        nome_completo = f"{nome} {sobrenome}"

        # Critério de busca (por parte do nome ou CPF exato)
        if (pesquisa in nome_completo.lower()) or (pesquisa == cpf):
            print("\n✅ Usuário Localizado com Sucesso!")
            print(f"👤 Nome Completo : {nome_completo}")
            print(f"🪪 CPF           : {cpf}")
            print(f"📧 E-mail        : {email}")
            print(f"📞 Telefone      : {telefone}")
            print(f"📅 Nascimento    : {nascimento}")
            print(f"⚡ Status        : {status}")
            print(f"📍 Endereço      : {endereco}")
            print(f"📝 Observação    : {observacao}")
            print("-" * 45)
            encontrado = True

    if not encontrado:
        print("\n🔍 Registro não encontrado na base de dados atual.")


if __name__ == "__main__":
    consultar_usuario()